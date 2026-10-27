"""Acceptance tests AC-1 to AC-29 of docs/design/convert-any-format.md.

They pin down what a user can observe: a document of any supported format turns
into markdown beside the original, the markdown says which bytes it came from so
an edited original is noticed, a hand-written file is never overwritten without
``--force``, and nothing leaves the machine. ``list-unconverted`` must keep
working in a core-only install, which is why several tests block the ``[ingest]``
modules from importing rather than uninstalling them.

The design does not name the dict that maps each markitdown suffix to the modules
that back it, so these tests fix the name ``ingest.MARKITDOWN_MODULES``. In
``list-unconverted`` output, the groups ``missing``, ``stale``, ``legacy``,
``unsupported`` and ``conflicting`` list the original (the path to pass to
``convert``), and ``orphaned`` lists the markdown file whose original is gone.
"""

# ruff: noqa: ANN202, S101, D100, D101, D102, D103, ANN001, ANN201, PLR2004, SLF001, INP001, RUF100, S603, S607, PLC0415, PLR0913, FBT003

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
import re
import socket
import subprocess
import sys
from pathlib import Path

import pytest
from conftest import GitBundle, needs_git
from typer.testing import CliRunner

from okf_kb import extras, frontmatter, ingest, okf

REPO_ROOT = Path(__file__).parents[2]
PLUGINS = REPO_ROOT / "plugins"

#: Every module the [ingest] extra brings, for simulating a core-only install.
INGEST_MODULES = (
    "markitdown",
    "fitz",
    "bs4",
    "markdownify",
    "requests",
    "mammoth",
    "pptx",
    "pandas",
    "openpyxl",
    "xlrd",
    "olefile",
)

LISTING_KEYS = {
    "missing",
    "stale",
    "unsupported",
    "orphaned",
    "legacy",
    "conflicting",
}


#: Runs ``list-unconverted`` in a fresh interpreter whose ``[ingest]`` modules
#: cannot be imported, which is what a core-only install looks like.
_LISTING_SCRIPT = """
import os, sys
root = sys.argv[1]
for name in sys.argv[2:]:
    sys.modules[name] = None
from typer.testing import CliRunner
from okf_kb import ingest
os.chdir(root)
result = CliRunner().invoke(ingest.app, ["list-unconverted", "--json-output"])
print(result.exit_code)
print(result.stdout)
"""


# -- fixtures built in code ----------------------------------------------------


def _png(size: int) -> bytes:
    import fitz

    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, size, size), False)
    pix.set_rect(pix.irect, (200, 40, 40))
    return pix.tobytes("png")


def _docx(
    path: Path,
    heading: str | None = "Quarterly report",
    paragraph: str = "The body paragraph.",
    image_sizes: tuple[int, ...] = (128,),
) -> Path:
    import docx
    from docx.shared import Inches

    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    if heading:
        document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    for size in image_sizes:
        document.add_picture(io.BytesIO(_png(size)), width=Inches(1))
    document.save(str(path))
    return path


def _pptx(path: Path, titles: tuple[str, ...] = ("Alpha", "Bravo", "Charlie")) -> Path:
    from pptx import Presentation

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    for title in titles:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
    presentation.save(str(path))
    return path


def _xlsx(path: Path) -> Path:
    import openpyxl

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Cost"])
    sheet.append(["Licences", 1234])
    workbook.save(path)
    return path


def _pdf(path: Path) -> Path:
    import fitz

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Figure 1 shows the architecture.")
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 128, 128), False)
    pix.set_rect(pix.irect, (40, 120, 200))
    page.insert_image(fitz.Rect(72, 100, 172, 200), pixmap=pix)
    doc.save(path)
    return path


def _bundle(tmp_path: Path, monkeypatch) -> Path:
    (tmp_path / "okf.toml").write_text("", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _invoke(*args: str | Path):
    return CliRunner().invoke(ingest.app, [str(a) for a in args])


def _read(path: Path) -> tuple[dict, str]:
    return frontmatter.parse(path.read_text(encoding="utf-8"))


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_md(path: Path, body: str = "# Converted\n", **meta: object) -> Path:
    path.write_text(frontmatter.dumps(meta, body), encoding="utf-8")
    return path


def _listing() -> dict[str, list[str]]:
    result = _invoke("list-unconverted", "--json-output")
    assert result.exit_code == 0, result.output
    return json.loads(result.stdout)


def _run_main(monkeypatch, capsys, *args: str | Path) -> tuple[int | str | None, str]:
    monkeypatch.setattr(sys, "argv", ["kb-ingest", *[str(a) for a in args]])
    with pytest.raises(SystemExit) as excinfo:
        ingest.main()
    return excinfo.value.code, capsys.readouterr().err


# -- AC-1 to AC-3: converting each format --------------------------------------


def test_ac_1_docx_converts_text_and_saves_the_image(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")

    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    md_path = source.with_suffix(".md")
    _, body = _read(md_path)
    assert "Quarterly report" in body
    assert "The body paragraph." in body
    images = list((root / "raw" / "images" / "report").iterdir())
    assert len(images) == 1
    links = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", body)
    assert len(links) == 1
    assert not links[0].startswith("data:")
    assert (md_path.parent / links[0]).resolve() == images[0].resolve()


def test_ac_2_pptx_slide_text_appears_in_order(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _pptx(root / "raw" / "documents" / "deck.pptx")

    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    _, body = _read(source.with_suffix(".md"))
    positions = [body.index(title) for title in ("Alpha", "Bravo", "Charlie")]
    assert positions == sorted(positions)


def test_ac_3_xlsx_converts_sheet_name_and_cells(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _xlsx(root / "raw" / "documents" / "budget.xlsx")

    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    _, body = _read(source.with_suffix(".md"))
    assert "Budget" in body
    assert "Licences" in body
    assert "1234" in body


# -- AC-4, AC-5: one implementation per conversion -----------------------------


def test_ac_4_pdf_body_matches_extract_pdf(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    legacy = _pdf(root / "raw" / "a" / "doc.pdf")
    converted = root / "raw" / "b" / "doc.pdf"
    converted.parent.mkdir(parents=True)
    converted.write_bytes(legacy.read_bytes())

    assert _invoke("extract-pdf", legacy).exit_code == 0
    result = _invoke("convert", converted)

    assert result.exit_code == 0, result.output
    meta, body = _read(converted.with_suffix(".md"))
    assert body == legacy.with_suffix(".md").read_text(encoding="utf-8")
    assert meta["converter"] == "pymupdf"


def test_ac_5_html_body_matches_html_to_md(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    html = (
        "<html><head><style>p{}</style></head><body><h1>Title</h1>"
        "<p>Some <b>bold</b> text</p><script>var a;</script></body></html>"
    )
    legacy = root / "raw" / "a" / "page.html"
    converted = root / "raw" / "b" / "page.html"
    for path in (legacy, converted):
        path.parent.mkdir(parents=True)
        path.write_text(html, encoding="utf-8")

    assert _invoke("html-to-md", legacy).exit_code == 0
    result = _invoke("convert", converted)

    assert result.exit_code == 0, result.output
    meta, body = _read(converted.with_suffix(".md"))
    assert body == legacy.with_suffix(".md").read_text(encoding="utf-8")
    assert meta["converter"] == "markdownify"


# -- AC-6 to AC-9: provenance and overwrite safety -----------------------------


def test_ac_6_converted_file_records_where_it_came_from(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")

    assert _invoke("convert", source).exit_code == 0

    meta, _ = _read(source.with_suffix(".md"))
    assert meta["converted_from"] == "report.docx"
    assert meta["converted_sha256"] == _sha(source)
    assert isinstance(meta["converted_sha256"], str)
    assert meta["converter"] == "markitdown"
    stamp = meta["converted"]
    if isinstance(stamp, str):
        stamp = dt.datetime.fromisoformat(stamp)
    assert isinstance(stamp, dt.datetime)
    for key in ("type", "title", "source_type", "date_added"):
        assert key in meta, key
    assert "author" not in meta
    assert "description" not in meta


def test_ac_7_unchanged_original_is_left_alone(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")
    assert _invoke("convert", source).exit_code == 0
    md_path = source.with_suffix(".md")
    touched = [md_path, *(root / "raw" / "images" / "report").iterdir()]
    old = 1_700_000_000_000_000_000
    for path in touched:
        os.utime(path, ns=(old, old))

    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    assert "up to date" in result.output.lower()
    assert [path.stat().st_mtime_ns for path in touched] == [old] * len(touched)


def test_ac_8_changed_original_is_stale_and_reconverted(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")
    assert _invoke("convert", source).exit_code == 0
    _docx(source, paragraph="A rewritten paragraph.")

    assert _listing()["stale"] == ["raw/documents/report.docx"]
    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    meta, body = _read(source.with_suffix(".md"))
    assert "A rewritten paragraph." in body
    assert meta["converted_sha256"] == _sha(source)
    listing = _listing()
    assert listing["stale"] == []
    assert listing["missing"] == []


def test_ac_9_hand_written_markdown_needs_force(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")
    md_path = source.with_suffix(".md")
    md_path.write_text("# My own notes\n", encoding="utf-8")
    before = md_path.read_bytes()

    refused = _invoke("convert", source)

    assert refused.exit_code == 1
    assert md_path.read_bytes() == before

    forced = _invoke("convert", source, "--force")

    assert forced.exit_code == 0, forced.output
    assert "The body paragraph." in md_path.read_text(encoding="utf-8")


# -- AC-10, AC-11: refusals ----------------------------------------------------


def test_ac_10_refuses_markdown_and_unsupported_suffixes(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    note = root / "raw" / "note.md"
    note.parent.mkdir(parents=True)
    note.write_text("# Note\n", encoding="utf-8")
    legacy = root / "raw" / "old.doc"
    legacy.write_bytes(b"not really a document")

    as_markdown = _invoke("convert", note)
    as_doc = _invoke("convert", legacy)

    assert as_markdown.exit_code == 1
    assert "already markdown" in as_markdown.output.lower()
    assert as_doc.exit_code == 1
    for suffix in (".docx", ".pptx", ".xlsx", ".pdf"):
        assert suffix in as_doc.output


def test_ac_11_two_originals_with_one_stem_conflict(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    pdf = _pdf(root / "raw" / "documents" / "a.pdf")
    assert _invoke("convert", pdf).exit_code == 0
    md_path = pdf.with_suffix(".md")
    before = md_path.read_bytes()
    docx_file = _docx(root / "raw" / "documents" / "a.docx")

    result = _invoke("convert", docx_file)

    assert result.exit_code == 1
    assert md_path.read_bytes() == before
    listing = _listing()
    assert listing["conflicting"] == ["raw/documents/a.docx"]
    assert "raw/documents/a.pdf" not in listing["missing"]


# -- AC-12, AC-13, AC-17: the listing ------------------------------------------


@needs_git
def test_ac_12_listing_groups_every_state(tmp_path, monkeypatch):
    bundle = GitBundle(tmp_path)
    monkeypatch.chdir(tmp_path)
    raw = tmp_path / "raw"
    docs = raw / "documents"

    _docx(docs / "fresh.docx")
    assert _invoke("convert", docs / "fresh.docx").exit_code == 0
    _docx(docs / "missing.docx")
    _docx(docs / "stale.docx")
    _write_md(
        docs / "stale.md",
        converted_from="stale.docx",
        converted_sha256="0" * 64,
    )
    _docx(docs / "legacy.docx")
    (docs / "legacy.md").write_text("# Hand written\n", encoding="utf-8")
    _write_md(docs / "gone.md", converted_from="gone.docx", converted_sha256="1" * 64)
    (docs / "old.doc").write_bytes(b"legacy format")
    (docs / "c.pdf").write_bytes(b"%PDF-1.4 placeholder")
    _write_md(
        docs / "c.md",
        converted_from="c.pdf",
        converted_sha256=_sha(docs / "c.pdf"),
    )
    _docx(docs / "c.docx")
    _docx(docs / "deleted.docx")
    _write_md(
        docs / "deleted.md",
        converted_from="deleted.docx",
        converted_sha256=_sha(docs / "deleted.docx"),
    )
    _docx(raw / "handwritten" / "scan.docx")
    _docx(raw / "images" / "x" / "inner.docx")
    _docx(docs / ".hidden.docx")
    bundle.commit("add raw documents")
    (docs / "deleted.docx").unlink()

    listing = _listing()

    assert set(listing) == LISTING_KEYS
    assert {key: sorted(value) for key, value in listing.items()} == {
        "missing": ["raw/documents/missing.docx"],
        "stale": ["raw/documents/stale.docx"],
        "unsupported": ["raw/documents/old.doc"],
        "orphaned": ["raw/documents/deleted.md", "raw/documents/gone.md"],
        "legacy": ["raw/documents/legacy.docx"],
        "conflicting": ["raw/documents/c.docx"],
    }


def test_ac_13_human_output(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    (root / "raw").mkdir()

    quiet = _invoke("list-unconverted")

    assert quiet.exit_code == 0
    lines = [line for line in quiet.stdout.splitlines() if line.strip()]
    assert len(lines) == 1
    assert "converted" in lines[0].lower()

    _docx(root / "raw" / "documents" / "todo.docx")
    (root / "raw" / "documents" / "old.doc").write_bytes(b"legacy format")

    busy = _invoke("list-unconverted")

    assert busy.exit_code == 0
    out = busy.stdout.lower()
    assert "missing" in out
    assert "unsupported" in out
    assert "todo.docx" in out
    assert "old.doc" in out


@needs_git
def test_ac_17_gitignored_files_are_skipped(tmp_path, monkeypatch):
    bundle = GitBundle(tmp_path)
    monkeypatch.chdir(tmp_path)
    bundle.write(".gitignore", "raw/ignored/\n")
    _docx(tmp_path / "raw" / "ignored" / "secret.docx")
    _docx(tmp_path / "raw" / "documents" / "new.docx")

    listing = _listing()

    assert listing["missing"] == ["raw/documents/new.docx"]
    everything = [path for group in listing.values() for path in group]
    assert not any("ignored" in path for path in everything)


# -- AC-14, AC-22, AC-24: a core-only install ----------------------------------


def test_ac_14_listing_works_without_the_extra_and_convert_says_how_to_fix(
    tmp_path,
    monkeypatch,
    capsys,
):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")
    for name in INGEST_MODULES:
        monkeypatch.setitem(sys.modules, name, None)

    listing = _invoke("list-unconverted", "--json-output")

    assert listing.exit_code == 0, listing.output
    assert json.loads(listing.stdout)["missing"] == ["raw/documents/report.docx"]
    code, err = _run_main(monkeypatch, capsys, "convert", source)
    assert code == 1
    assert "[ingest]" in err
    assert "uv" in err
    assert "Traceback" not in err


@pytest.mark.parametrize(
    "case",
    [
        (_docx, "report.docx", "mammoth"),
        (_pptx, "deck.pptx", "pptx"),
        (_xlsx, "budget.xlsx", "pandas"),
    ],
    ids=["docx-mammoth", "pptx-pptx", "xlsx-pandas"],
)
def test_ac_22_missing_sub_dependency_is_reported_not_raised(
    case,
    tmp_path,
    monkeypatch,
    capsys,
):
    build, name, module = case
    root = _bundle(tmp_path, monkeypatch)
    source = build(root / "raw" / "documents" / name)
    monkeypatch.setitem(sys.modules, module, None)

    code, err = _run_main(monkeypatch, capsys, "convert", source)

    assert code == 1
    assert "[ingest]" in err
    assert "uv" in err
    assert "Traceback" not in err


def test_ac_24_ingest_imports_and_lists_with_no_extra_installed(tmp_path):
    (tmp_path / "okf.toml").write_text("", encoding="utf-8")
    (tmp_path / "raw" / "documents").mkdir(parents=True)
    (tmp_path / "raw" / "documents" / "a.docx").write_bytes(b"placeholder")
    done = subprocess.run(
        [sys.executable, "-c", _LISTING_SCRIPT, str(tmp_path), *INGEST_MODULES],
        capture_output=True,
        text=True,
        check=False,
        cwd=tmp_path,
    )

    assert done.returncode == 0, done.stderr
    code, _, payload = done.stdout.partition("\n")
    assert code == "0"
    assert json.loads(payload)["missing"] == ["raw/documents/a.docx"]


# -- AC-15: nothing leaves the machine -----------------------------------------


@pytest.mark.parametrize(
    ("build", "name"),
    [(_docx, "report.docx"), (_pptx, "deck.pptx"), (_xlsx, "budget.xlsx")],
)
def test_ac_15_conversion_makes_no_network_connection(
    build,
    name,
    tmp_path,
    monkeypatch,
):
    root = _bundle(tmp_path, monkeypatch)
    source = build(root / "raw" / "documents" / name)
    attempts: list[object] = []

    def refuse(*args: object, **_kwargs: object):
        attempts.append(args)
        msg = "network access is blocked in this test"
        raise OSError(msg)

    monkeypatch.setattr(socket.socket, "connect", refuse)
    monkeypatch.setattr(socket.socket, "connect_ex", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(socket, "getaddrinfo", refuse)

    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    assert attempts == []


# -- AC-16: small images are skipped -------------------------------------------


def test_ac_16_icon_sized_images_are_skipped(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx", image_sizes=(16, 128))

    result = _invoke("convert", source)

    assert result.exit_code == 0, result.output
    assert len(list((root / "raw" / "images" / "report").iterdir())) == 1
    _, body = _read(source.with_suffix(".md"))
    assert len(re.findall(r"!\[[^\]]*\]\(", body)) == 1


# -- AC-19, AC-23: declarations ------------------------------------------------


def test_ac_19_extras_know_about_markitdown(monkeypatch):
    assert "markitdown" in extras.EXTRA_MODULES["ingest"]
    monkeypatch.setitem(sys.modules, "markitdown", None)

    assert "markitdown" in extras.missing_modules("ingest")


def test_ac_23_every_markitdown_suffix_names_its_modules():
    assert set(ingest.MARKITDOWN_SUFFIXES) == {
        ".docx",
        ".pptx",
        ".xlsx",
        ".xls",
        ".epub",
        ".msg",
    }
    assert set(ingest.MARKITDOWN_MODULES) == set(ingest.MARKITDOWN_SUFFIXES)
    declared = set(extras.EXTRA_MODULES["ingest"])
    for suffix, modules in ingest.MARKITDOWN_MODULES.items():
        assert set(modules) <= declared, suffix


# -- AC-20, AC-28, AC-29: the skills -------------------------------------------


def test_ac_20_skills_call_the_new_commands():
    compile_skill = (PLUGINS / "kb/skills/compile/SKILL.md").read_text(
        encoding="utf-8",
    )
    ingest_skill = (PLUGINS / "kb-ingest/skills/ingest/SKILL.md").read_text(
        encoding="utf-8",
    )

    assert "kb-ingest list-unconverted" in compile_skill
    assert "kb-ingest convert" in compile_skill
    assert "kb-ingest convert" in ingest_skill


def test_ac_28_init_creates_the_documents_folder():
    skill = (PLUGINS / "kb/skills/init/SKILL.md").read_text(encoding="utf-8")
    template = (PLUGINS / "kb/skills/init/templates/CLAUDE.md").read_text(
        encoding="utf-8",
    )

    assert re.search(r"mkdir -p [^\n]*<raw>/documents", skill)
    assert "documents/" in template


def test_ac_29_adopt_offers_to_convert_loose_documents():
    text = (PLUGINS / "kb/skills/adopt/SKILL.md").read_text(encoding="utf-8")

    for suffix in (".docx", ".pptx", ".xlsx"):
        assert suffix in text
    assert "documents/" in text
    assert "kb-ingest convert" in text
    assert "kb-doctor" in text
    assert "uv tool install" not in text


# -- AC-25 to AC-27: the frontmatter a converted file carries ------------------


@needs_git
def test_ac_25_converted_files_carry_the_raw_identity_keys(tmp_path, monkeypatch):
    GitBundle(tmp_path)
    monkeypatch.chdir(tmp_path)
    raw = tmp_path / "raw"
    cases = [
        (
            _docx(raw / "papers" / "p.docx", heading="Paper heading"),
            "paper",
            "Paper heading",
        ),
        (_docx(raw / "documents" / "d.docx", heading=None), "document", "d"),
        (
            _docx(raw / "misc" / "m.docx", heading="Misc heading"),
            "document",
            "Misc heading",
        ),
    ]
    today = dt.datetime.now(dt.UTC).date()

    for source, expected_type, expected_title in cases:
        result = _invoke("convert", source)

        assert result.exit_code == 0, result.output
        meta, _ = _read(source.with_suffix(".md"))
        assert meta["type"] == expected_type, source.name
        assert meta["title"] == expected_title, source.name
        assert meta["source_type"] == "technical"
        assert frontmatter.as_date(meta["date_added"]) == today
        assert "author" not in meta
        assert "description" not in meta


def test_ac_26_regeneration_keeps_every_other_key(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    extra = {
        "compile": False,
        "tags": ["a", "b"],
        "description": "Written by hand.",
        "title": "Hand edited title",
    }
    source = _docx(root / "raw" / "documents" / "report.docx")
    assert _invoke("convert", source).exit_code == 0
    md_path = source.with_suffix(".md")
    meta, body = _read(md_path)
    old_hash = meta["converted_sha256"]
    md_path.write_text(frontmatter.dumps({**meta, **extra}, body), encoding="utf-8")
    _docx(source, paragraph="A rewritten paragraph.")

    assert _invoke("convert", source).exit_code == 0

    after, new_body = _read(md_path)
    for key, value in extra.items():
        assert after[key] == value, key
    assert after["converted_sha256"] != old_hash
    assert after["converted_sha256"] == _sha(source)
    assert "A rewritten paragraph." in new_body

    forced = _docx(root / "raw" / "documents" / "second.docx")
    forced.with_suffix(".md").write_text(
        frontmatter.dumps(extra, "# Hand written\n"),
        encoding="utf-8",
    )

    assert _invoke("convert", forced, "--force").exit_code == 0

    forced_meta, forced_body = _read(forced.with_suffix(".md"))
    for key, value in extra.items():
        assert forced_meta[key] == value, key
    assert forced_meta["converted_sha256"] == _sha(forced)
    assert "The body paragraph." in forced_body
    assert "Hand written" not in forced_body


def test_ac_27_migrating_a_converted_file_keeps_its_provenance(tmp_path, monkeypatch):
    root = _bundle(tmp_path, monkeypatch)
    source = _docx(root / "raw" / "documents" / "report.docx")
    assert _invoke("convert", source).exit_code == 0
    md_path = source.with_suffix(".md")
    document = frontmatter.load(md_path)

    migrated = okf.migrate_raw(
        document,
        "raw/documents/report.md",
        root,
        raw_name="raw",
    )

    for key in ("converted_from", "converted_sha256", "converter", "converted"):
        assert migrated.frontmatter[key] == document.frontmatter[key], key
    assert migrated.body == document.body
